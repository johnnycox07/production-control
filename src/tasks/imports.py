import os
import tempfile

import openpyxl

from src.celery_app import celery_app
from src.core.sync_database import SyncSessionLocal as SyncSession


@celery_app.task(bind=True, max_retries=1)
def import_batches_from_file(self, file_url: str):
    from src.data.models.batch import Batch
    from src.data.models.work_center import WorkCenter
    from src.storage.minio_service import MinIOService


    minio = MinIOService()
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        tmp_path = tmp.name

    object_name = file_url.split("/")[-1]
    minio.download_file("imports", object_name, tmp_path)

    wb = openpyxl.load_workbook(tmp_path)
    ws = wb.active

    total_rows = 0
    created = 0
    skipped = 0
    errors = []

    with SyncSession() as session:
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not any(row):
                continue

            total_rows += 1

            try:
                (
                    batch_number, batch_date, nomenclature, ekn_code,
                    work_center_name, work_center_identifier,
                    shift, team, shift_start, shift_end,
                ) = row[:10]

                work_center = (
                    session.query(WorkCenter)
                    .filter(WorkCenter.identifier == work_center_identifier)
                    .first()
                )
                if not work_center:
                    work_center = WorkCenter(
                        identifier=work_center_identifier,
                        name=work_center_name,
                    )
                    session.add(work_center)
                    session.flush()

                exists = (
                    session.query(Batch)
                    .filter(
                        Batch.batch_number == batch_number,
                        Batch.batch_date == batch_date,
                    )
                    .first()
                )
                if exists:
                    skipped += 1
                    errors.append({
                        "row": row_idx,
                        "error": "Duplicate batch number and date",
                    })
                    continue

                batch = Batch(
                    batch_number=batch_number,
                    batch_date=batch_date,
                    nomenclature=nomenclature,
                    ekn_code=ekn_code,
                    work_center_id=work_center.id,
                    shift=shift,
                    team=team,
                    shift_start=shift_start,
                    shift_end=shift_end,
                    task_description=f"Импорт: {nomenclature}",
                )
                session.add(batch)
                created += 1

                if total_rows % 10 == 0:
                    self.update_state(
                        state="PROGRESS",
                        meta={
                            "current": total_rows,
                            "total": ws.max_row - 1,
                            "created": created,
                            "skipped": skipped,
                        }
                    )

            except Exception as e:
                skipped += 1
                errors.append({"row": row_idx, "error": str(e)})
                continue

        session.commit()

    os.unlink(tmp_path)

    return {
        "success": True,
        "total_rows": total_rows,
        "created": created,
        "skipped": skipped,
        "errors": errors,
    }